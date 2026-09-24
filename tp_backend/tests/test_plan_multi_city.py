"""Planning a trip that covers two cities far apart: one shortlist, one day, one draft, one workbook.

Everything here is Helsinki + Porto over midsummer, where the two sunsets are an hour and a half
apart — close enough that both are plausible, far enough that using the wrong one is visible.
"""

from datetime import timedelta
from io import BytesIO

from conftest import HELSINKI, make_mention, make_place
from openpyxl import load_workbook
from sqlalchemy import select
from test_multi_city import PORTO, midsummer, two_cities

from libs.db import City, ItineraryItem, Place, Trip, TripPlace
from libs.db.enums import Confidence, TaskKind
from libs.routing import HoursHit
from tp_api.route_planning.export import OK_BG, WARN_BG
from tp_api.route_planning.service import shortlist
from tp_api.route_planning.utils import google_weekday
from tp_ingestions.plan import draft
from tp_ingestions.queue import ClaimedTask

SINGAPORE = "ChIJdZOLiiMR2jERnbSmdlpEFRI"


def two_city_trip(client, lookup, days=1, **kw):
    arrive = midsummer()
    return client.post("/initiate-plan", json=two_cities(
        lookup, arrive_date=arrive.isoformat(), arrive_time=None, depart_time=None,
        depart_date=(arrive + timedelta(days=days)).isoformat(), **kw)).json()["trip_id"]


def pin(client, trip, items, day=0):
    r = client.put(f"/trips/{trip}/itinerary", json={"days": [{"day_index": day, "items": items}]})
    assert r.status_code == 200, r.text
    return r.json()


def block(place_id, start_min, duration_min=60, **kw):
    return {"place_id": place_id, "start_min": start_min, "duration_min": duration_min} | kw


def open_until(hour, minute=0):
    return lambda ids: {i: HoursHit(place_id=i, periods=[
        {"open": {"day": d, "hour": 8, "minute": 0},
         "close": {"day": d, "hour": hour, "minute": minute}} for d in range(7)],
        weekday_descriptions=[], utc_offset_minutes=60) for i in ids}


def one_in_each_city(db, trip, hel_start, por_start, **por):
    make_place(db, city_id=HELSINKI, place_id="hel", name="Suomenlinna", lat=60.14, lon=24.98)
    make_mention(db, "hel", category="see")
    make_place(db, city_id=PORTO, place_id="por", name="Jardim do Morro",
               **{"lat": 41.15, "lon": -8.61} | por)
    make_mention(db, "por", category="see", source_ref="ref2")
    return [block("por", por_start), block("hel", hel_start)]


def test_the_ranking_ignores_which_city_a_place_is_in(client, db, lookup):
    trip = two_city_trip(client, lookup)
    make_place(db, city_id=HELSINKI, place_id="hel", name="Once named")
    make_mention(db, "hel")
    make_place(db, city_id=PORTO, place_id="por", name="Thrice named")
    for i in range(3):
        make_mention(db, "por", source_ref=f"por-{i}")

    places = client.get(f"/trips/{trip}/shortlist").json()["places"]

    assert [(p["name"], p["city_id"]) for p in places] == [("Thrice named", PORTO),
                                                           ("Once named", HELSINKI)]


def test_a_day_mixing_two_cities_draws_no_complaint(client, db, lookup, hours):
    trip = two_city_trip(client, lookup)
    hours["fn"] = open_until(23)

    pin(client, trip, one_in_each_city(db, trip, hel_start=13 * 60, por_start=12 * 60))
    body = client.post(f"/trips/{trip}/days/0/route", json={}).json()

    assert [b["place_id"] for b in body["blocks"]] == ["por", "hel"]
    assert body["warnings"] == []


def test_the_day_s_daylight_is_the_first_block_s_city_not_the_anchor_s(client, db, lookup, hours):
    trip = two_city_trip(client, lookup)
    hours["fn"] = open_until(23)
    pin(client, trip, one_in_each_city(db, trip, hel_start=13 * 60, por_start=12 * 60))

    body = client.post(f"/trips/{trip}/days/0/route", json={}).json()

    assert body["daylight"] == {"sunrise": "06:01", "sunset": "21:10"}, "Porto's"


def test_a_place_with_no_coordinates_falls_back_to_its_own_city_s_centre(client, db, lookup, hours):
    trip = two_city_trip(client, lookup)
    hours["fn"] = open_until(23)
    pin(client, trip, one_in_each_city(db, trip, hel_start=13 * 60, por_start=21 * 60 + 30,
                                       lat=None, lon=None))

    body = client.post(f"/trips/{trip}/days/0/route", json={}).json()

    dark = [w for w in body["warnings"] if w["code"] == "after_sunset"]
    assert [(w["place_id"], w["detail"]["sunset"]) for w in dark] == [("por", "21:10")]


def test_a_city_with_no_timezone_falls_back_to_the_offset_places_reported(client, db, lookup,
                                                                         hours):
    trip = two_city_trip(client, lookup)
    hours["fn"] = lambda ids: {i: HoursHit(place_id=i, periods=[], weekday_descriptions=[],
                                           utc_offset_minutes=180) for i in ids}
    make_place(db, city_id=PORTO, place_id="por", name="Jardim do Morro", lat=41.14, lon=-8.61)
    db.get(City, PORTO).timezone = None
    db.commit()
    pin(client, trip, [block("por", 12 * 60)])

    body = client.post(f"/trips/{trip}/days/0/route", json={}).json()

    assert body["daylight"] == {"sunrise": "08:01", "sunset": "23:10"}, "Porto at +180, not at UTC"


def test_a_place_and_city_with_no_coordinates_report_no_daylight(client, db, lookup, hours):
    trip = two_city_trip(client, lookup)
    hours["fn"] = open_until(23)
    make_place(db, city_id=PORTO, place_id="por", name="Jardim do Morro", lat=None, lon=None)
    make_mention(db, "por", category="see")
    porto = db.get(City, PORTO)
    porto.lat = porto.lon = None
    db.commit()
    pin(client, trip, [block("por", 22 * 60 + 30, 30)])

    body = client.post(f"/trips/{trip}/days/0/route", json={}).json()

    assert body["daylight"] is None
    assert [w["code"] for w in body["warnings"]] == []


def test_polar_day_is_no_daylight_rather_than_a_sunset(client, db, lookup, hours):
    trip = two_city_trip(client, lookup)
    hours["fn"] = open_until(23, 30)
    make_place(db, city_id=HELSINKI, place_id="pol", name="Longyearbyen", lat=78.9, lon=11.9)
    make_mention(db, "pol", category="see")
    pin(client, trip, [block("pol", 23 * 60, 30)])

    body = client.post(f"/trips/{trip}/days/0/route", json={}).json()

    assert body["daylight"] is None
    assert "after_sunset" not in [w["code"] for w in body["warnings"]]


def singapore_place(db, trip):
    db.add(City(city_id=SINGAPORE, name="Singapore", country="SG", timezone="Asia/Singapore",
                lat=1.35, lon=103.82))
    db.add(Place(place_id="sgp", city_id=SINGAPORE, name="Gardens by the Bay", lat=1.28, lon=103.86,
                 confidence=Confidence.HIGH, category="see"))
    db.flush()
    db.add(TripPlace(trip_id=trip, place_id="sgp"))
    db.commit()


def test_a_place_filed_under_an_uncovered_city_is_judged_where_it_actually_is(client, db, lookup,
                                                                             hours):
    trip = two_city_trip(client, lookup)
    hours["fn"] = open_until(23)
    singapore_place(db, trip)
    pin(client, trip, [block("sgp", 21 * 60 + 30)])

    body = client.post(f"/trips/{trip}/days/0/route", json={}).json()

    assert body["daylight"] == {"sunrise": "07:00", "sunset": "19:11"}
    assert [(w["code"], w["detail"]["sunset"]) for w in body["warnings"]] == \
        [("after_sunset", "19:11")]


def test_the_draft_prompt_names_a_city_the_trip_does_not_cover(client, db, lookup):
    trip = two_city_trip(client, lookup)
    singapore_place(db, trip)

    prompt = draft.render(db, db.get(Trip, trip), shortlist(db, trip, 40, 0, None).places, [0])

    row = next(ln for ln in prompt.splitlines() if "Gardens by the Bay" in ln)
    assert "Singapore" in row and "0.0km" not in row, row


def test_more_stored_stops_than_a_day_takes_is_a_422(client, db, lookup):
    trip = two_city_trip(client, lookup)
    for i in range(26):
        make_place(db, city_id=PORTO if i % 2 else HELSINKI, place_id=f"p{i}", name=f"Place {i}")
        db.add(ItineraryItem(trip_id=trip, place_id=f"p{i}", day_index=0,
                             start_min=8 * 60 + i * 30, duration_min=30))
    db.commit()

    assert client.post(f"/trips/{trip}/days/0/route", json={}).status_code == 422


def workbook(client, trip):
    r = client.get(f"/trips/{trip}/export.xlsx")
    assert r.status_code == 200, r.text
    return load_workbook(BytesIO(r.content))["Itinerary"]


def a_dark_porto_and_a_light_helsinki(client, db, lookup, hours, **kw):
    trip = two_city_trip(client, lookup, days=1)
    hours["fn"] = open_until(23, 30)
    items = one_in_each_city(db, trip, hel_start=22 * 60, por_start=21 * 60 + 30)
    pin(client, trip, [items[0], items[1] | kw])
    return trip


def test_the_export_bands_each_day_with_its_first_block_s_own_daylight(client, db, lookup, hours):
    trip = a_dark_porto_and_a_light_helsinki(client, db, lookup, hours)

    ws = workbook(client, trip)

    assert ws.cell(row=4, column=1).value == \
        f"Day 1 · {midsummer():%a %d %b} · 2 blocks · daylight 06:01–21:10"
    assert ws.cell(row=7, column=1).value.endswith("no blocks"), "an empty day bands with no light"


def test_the_export_warns_only_the_block_that_is_actually_dark(client, db, lookup, hours):
    trip = a_dark_porto_and_a_light_helsinki(client, db, lookup, hours)

    ws = workbook(client, trip)

    assert [ws.cell(row=r, column=6).value for r in (5, 6)] == ["Jardim do Morro", "Suomenlinna"]
    assert ws.cell(row=5, column=8).value == "dark by 21:30 — sunset 21:10"
    assert ws.cell(row=5, column=8).fill.fgColor.rgb.endswith(WARN_BG)
    assert ws.cell(row=6, column=8).value == "ok"
    assert ws.cell(row=6, column=8).fill.fgColor.rgb.endswith(OK_BG)


def test_the_ref_column_marks_only_the_block_that_carries_a_link(client, db, lookup, hours):
    trip = a_dark_porto_and_a_light_helsinki(client, db, lookup, hours,
                                             reference_url="https://www.airbnb.com/rooms/12345")

    ws = workbook(client, trip)

    assert [c.value for c in ws[3]][-1] == "Ref"
    assert ws.cell(row=5, column=9).value is None, "Porto has no link of its own"
    assert ws.cell(row=6, column=9).hyperlink.target == "https://www.airbnb.com/rooms/12345"


def drivable(client, db, lookup, monkeypatch, reply, hours_for):
    trip = client.post("/initiate-plan", json=two_cities(lookup)).json()["trip_id"]
    make_place(db, city_id=HELSINKI, place_id="hel", name="Löyly")
    for i in range(2):
        make_mention(db, "hel", category="do", source_ref=f"hel-{i}")
    make_place(db, city_id=PORTO, place_id="por", name="Bolhão")
    make_mention(db, "por", category="eat", source_ref="por-0")

    calls = []
    monkeypatch.setattr(draft, "generate", lambda prompt, text: calls.append(text) or reply)
    monkeypatch.setattr(draft, "fetch_hours", hours_for)
    return trip, calls


def periods(days, open_hour=8, close_hour=23):
    return [{"open": {"day": d, "hour": open_hour, "minute": 0},
             "close": {"day": d, "hour": close_hour, "minute": 0}} for d in days]


def hours_of(per_place):
    return lambda ids: {i: HoursHit(place_id=i, periods=per_place(i), weekday_descriptions=[],
                                    utc_offset_minutes=180) for i in ids}


def task(trip):
    return ClaimedTask(task_id=1, run_id="r-plan", kind=TaskKind.ROUTE_PLAN, source=None,
                       payload={"trip_id": trip}, attempts=1, max_attempts=3)


BOTH = {"days": [{"day": 0, "picks": [{"index": 0, "start_min": 900, "duration_min": 60},
                                      {"index": 1, "start_min": 1080, "duration_min": 60}]}]}


def test_the_draft_lands_blocks_from_both_cities(client, db, lookup, monkeypatch):
    trip, calls = drivable(client, db, lookup, monkeypatch, BOTH,
                           hours_of(lambda i: periods(range(7))))

    assert draft.run(db, task(trip)) == {"days": 1, "blocks": 2, "rounds": 1, "unresolved": 0}
    assert len(calls) == 1
    assert set(db.scalars(select(ItineraryItem.place_id))) == {"hel", "por"}
    assert "Bolhão" in calls[0] and "Löyly" in calls[0]


def test_the_draft_drops_a_place_proven_closed_all_day(client, db, lookup, monkeypatch):
    trip, calls = drivable(client, db, lookup, monkeypatch, BOTH, hours_of(
        lambda i: periods(range(7))))
    tomorrow = (google_weekday(db.get(Trip, trip).arrive_date) + 1) % 7
    monkeypatch.setattr(draft, "fetch_hours", hours_of(
        lambda i: periods([tomorrow] if i == "por" else range(7))))

    result = draft.run(db, task(trip))

    assert result == {"days": 1, "blocks": 1, "rounds": 2, "unresolved": 0}
    assert len(calls) == 2, "one retry, with the closed place named in the feedback"
    assert db.scalars(select(ItineraryItem.place_id)).all() == ["hel"]


def test_a_draft_with_nothing_shortlisted_spends_no_model_call(client, db, lookup, monkeypatch):
    trip = client.post("/initiate-plan", json=two_cities(lookup)).json()["trip_id"]
    monkeypatch.setattr(draft, "generate", lambda *a: 1 / 0)

    assert draft.run(db, task(trip)) == {"skipped": "nothing resolved for this trip's cities yet"}
