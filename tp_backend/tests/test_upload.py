"""Upload new trip: only this server's own export comes back, edited in Excel or not."""

import base64
import re
from datetime import date, timedelta
from io import BytesIO

import pytest
from conftest import make_place, plan_body
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from libs.db import ItineraryItem, UserTrip
from libs.db.enums import TripRole
from tp_api.route_planning.export import META_SHEET, until
from tp_api.uploads import _end


def planned_trip(client, db):
    trip = client.post("/initiate-plan", json=plan_body(name="Nordic")).json()["trip_id"]
    make_place(db, place_id="p1", name="Löyly")
    client.put(f"/trips/{trip}/itinerary", json={"days": [
        {"day_index": 0, "items": [{"place_id": "p1", "start_min": 900, "duration_min": 90}]},
        {"day_index": 1, "items": [{"kind": "custom", "block_id": "b1", "title": "Night train",
                                    "description": "Car 4", "start_min": 1350,
                                    "duration_min": 460}]}]})
    return trip


def exported(client, trip) -> bytes:
    return client.get(f"/trips/{trip}/export.xlsx").content


def edited(content: bytes, change) -> bytes:
    wb = load_workbook(BytesIO(content))
    change(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, content: bytes, path="/uploads", **kw):
    return client.post(path, json={"file": base64.b64encode(content).decode()} | kw)


def items(db, trip):
    return [(i.day_index, i.kind, i.place_id, i.title, i.start_min, i.duration_min)
            for i in db.scalars(select(ItineraryItem).where(ItineraryItem.trip_id == trip)
                                .order_by(ItineraryItem.day_index))]


def test_an_export_comes_back_as_a_new_trip_the_uploader_owns(client, db):
    source = planned_trip(client, db)
    r = upload(client, exported(client, source))
    assert r.status_code == 200, r.text
    trip = r.json()["trip_id"]

    assert trip != source
    assert r.json()["name"] == "Nordic"
    assert items(db, trip) == [(0, "place", "p1", None, 900, 90),
                               (1, "custom", None, "Night train", 1350, 460)]
    role = db.scalar(select(UserTrip.role).where(UserTrip.trip_id == trip))
    assert role == TripRole.OWNER
    assert len(items(db, source)) == 2


def test_the_preview_writes_nothing_and_shows_the_days(client, db):
    source = planned_trip(client, db)
    body = upload(client, exported(client, source), "/uploads/preview").json()

    assert body["cities"] == ["Helsinki"]
    assert [[(b["name"], b["start"], b["end"]) for b in d["blocks"]] for d in body["days"]] == [
        [("Löyly", "15:00", "16:30")], [("Night train", "22:30", "Day 3 06:10")]]
    assert body["skipped"] == []
    assert len(client.get("/trips").json()) == 1


def test_a_new_start_date_moves_every_day_and_keeps_the_length(client, db):
    source = planned_trip(client, db)
    later = date.fromisoformat(plan_body()["arrive_date"]) + timedelta(days=60)
    trip = upload(client, exported(client, source), arrive_date=later.isoformat()).json()

    assert trip["arrive_date"] == later.isoformat()
    assert trip["depart_date"] == (later + timedelta(days=3)).isoformat()


def test_edits_made_in_excel_come_through(client, db):
    source = planned_trip(client, db)

    def change(wb):
        ws = wb["Itinerary"]
        ws["D5"] = "16:00"
        ws["F7"] = "Sleeper to Rovaniemi"
        ws["D11"], ws["E11"], ws["F11"] = "10:00", "11:15", "Husky sled"

    trip = upload(client, edited(exported(client, source), change)).json()["trip_id"]
    assert items(db, trip) == [(0, "place", "p1", None, 960, 30),
                               (1, "custom", None, "Sleeper to Rovaniemi", 1350, 460),
                               (3, "custom", None, "Husky sled", 600, 75)]


def test_a_place_on_two_days_comes_back_on_both_with_its_block_ids(client, db):
    trip = client.post("/initiate-plan", json=plan_body()).json()["trip_id"]
    make_place(db, place_id="p1", name="Hotel Kämp")
    client.put(f"/trips/{trip}/itinerary", json={"days": [
        {"day_index": d, "items": [{"place_id": "p1", "block_id": f"night-{d}",
                                    "start_min": 1200, "duration_min": 60}]} for d in (0, 1)]})

    copy = upload(client, exported(client, trip)).json()["trip_id"]
    rows = db.execute(select(ItineraryItem.day_index, ItineraryItem.place_id, ItineraryItem.block_id)
                      .where(ItineraryItem.trip_id == copy).order_by(ItineraryItem.day_index)).all()
    assert [tuple(r) for r in rows] == [(0, "p1", "night-0"), (1, "p1", "night-1")]


def test_an_export_without_place_block_ids_still_reads(client, db):
    source = planned_trip(client, db)

    def change(wb):
        meta = wb[META_SHEET]
        meta["A2"], n = re.subn(r'"block_id":"[^"]+","kind":"place"',
                                '"block_id":null,"kind":"place"', meta["A2"].value)
        assert n == 1

    trip = upload(client, edited(exported(client, source), change)).json()["trip_id"]
    assert items(db, trip)[0] == (0, "place", "p1", None, 900, 90)


def test_a_renamed_place_becomes_a_custom_block(client, db):
    source = planned_trip(client, db)
    trip = upload(client, edited(exported(client, source),
                                 lambda wb: wb["Itinerary"].__setitem__("F5", "Sauna night"))
                  ).json()["trip_id"]
    assert items(db, trip)[0] == (0, "custom", None, "Sauna night", 900, 90)


def test_an_unreadable_row_is_skipped_and_named(client, db):
    source = planned_trip(client, db)
    body = upload(client, edited(exported(client, source),
                                 lambda wb: wb["Itinerary"].__setitem__("D5", "lunch")),
                  "/uploads/preview").json()
    assert body["skipped"] == [{"row": 5, "reason": "Start or End is not a time like 14:30"}]


@pytest.mark.parametrize("content, reason", [
    (b"not a spreadsheet", "That is not an .xlsx file."),
    (None, "This file was not exported from Trip Planner."),
])
def test_anything_else_is_refused(client, content, reason):
    if content is None:
        buf = BytesIO()
        Workbook().save(buf)
        content = buf.getvalue()
    r = upload(client, content)
    assert r.status_code == 422
    assert r.json()["detail"] == reason


def test_a_damaged_hidden_sheet_is_refused(client, db):
    source = planned_trip(client, db)
    broken = edited(exported(client, source), lambda wb: wb[META_SHEET].__setitem__("A2", "{oops"))
    r = upload(client, broken)
    assert r.status_code == 422
    assert "damaged" in r.json()["detail"]


def test_damaged_row_data_skips_only_that_row(client, db):
    source = planned_trip(client, db)

    def change(wb):
        meta = wb[META_SHEET]
        meta["A2"] = meta["A2"].value.replace('"reference_url":null', '"reference_url":"nope"', 1)

    body = upload(client, edited(exported(client, source), change), "/uploads/preview").json()
    assert body["skipped"] == [{"row": 5, "reason": "its hidden Trip Planner data is damaged"}]
    assert [d["day_index"] for d in body["days"]] == [1]


@pytest.mark.parametrize("day, end, text, back", [
    (0, 16 * 60 + 30, "16:30", (0, 990)),
    (0, 24 * 60, "24:00", (0, 1440)),
    (1, 1350 + 460, "Day 3 06:10", (2, 370)),
    (1, 2 * 24 * 60, "Day 3 24:00", (2, 1440)),
])
def test_an_end_is_written_the_way_it_is_read(day, end, text, back):
    assert until(day, end) == text
    assert _end(text, day) == back
