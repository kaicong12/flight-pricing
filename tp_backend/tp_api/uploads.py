"""Upload new trip: read back a workbook this server exported, and refuse everything else.

The hidden sheet says which trip the file came from and what each row was; the visible Itinerary
sheet says what the days are now, so edits made in Excel come through.
"""

import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from libs.db import Place, Trip
from libs.db.enums import BlockKind
from libs.settings import settings
from tp_api.route_planning.export import MARKER, META_SHEET, ROW_KEY, VERSION
from tp_api.route_planning.schemas import DESCRIPTION_MAX, TITLE_MAX, ItemIn
from tp_api.route_planning.utils import available_window, day_count, pieces

DAY = 24 * 60
BAND = re.compile(r"^Day (\d+) ·")
CLOCK = re.compile(r"^(\d{1,2}):(\d{2})(?::\d{2})?$")
LATER = re.compile(r"^Day\s+(\d+)\s+(\d{1,2}:\d{2}(?::\d{2})?)$")


DAMAGED = "This file's hidden Trip Planner sheet is damaged. Export the trip again."


class NotOurExport(ValueError):
    pass


@dataclass
class Upload:
    trip: dict
    days: dict[int, list[ItemIn]] = field(default_factory=dict)
    skipped: list[tuple[int, str]] = field(default_factory=list)

    @property
    def place_ids(self) -> list[str]:
        return [i.place_id for items in self.days.values() for i in items if i.place_id]


def _clock(value) -> int | None:
    """Minutes past midnight from whatever Excel left in the cell. "24:00" is a day's end."""
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)) and 0 <= value <= 1:
        return round(value * DAY)
    m = CLOCK.match(str(value or "").strip())
    if not m or int(m[2]) > 59 or int(m[1]) * 60 + int(m[2]) > DAY:
        return None
    return int(m[1]) * 60 + int(m[2])


def _end(value, day: int) -> tuple[int, int] | None:
    m = LATER.match(str(value or "").strip())
    if m:
        at = _clock(m[2])
        return None if at is None else (int(m[1]) - 1, at)
    at = _clock(value)
    return None if at is None else (day, at)


def _meta(wb) -> dict:
    if META_SHEET not in wb.sheetnames:
        raise NotOurExport("This file was not exported from Trip Planner.")
    ws = wb[META_SHEET]
    if ws.cell(row=1, column=1).value != MARKER:
        raise NotOurExport("This file was not exported from Trip Planner.")
    if ws.cell(row=1, column=2).value != VERSION:
        raise NotOurExport("This export is from a different version of Trip Planner. "
                           "Export the trip again.")
    payload = "".join(str(ws.cell(row=r, column=1).value or "") for r in range(2, ws.max_row + 1))
    try:
        meta = json.loads(payload)
        meta["trip"]["cities"][0]["city_id"]
        if not all(isinstance(v, dict) for v in meta["rows"].values()):
            raise TypeError("a row is not an object")
    except (ValueError, KeyError, IndexError, TypeError, AttributeError) as e:
        raise NotOurExport(DAMAGED) from e
    return meta


def read(db: Session, content: bytes, arrive_date: date | None = None) -> Upload:
    """The trip the file describes, moved to `arrive_date` if given, and the rows that fit it."""
    try:
        wb = load_workbook(BytesIO(content))
    except Exception as e:
        raise NotOurExport("That is not an .xlsx file.") from e
    meta = _meta(wb)
    if "Itinerary" not in wb.sheetnames:
        raise NotOurExport("The Itinerary sheet is missing.")

    info = meta["trip"]
    try:
        start = date.fromisoformat(info["arrive_date"])
        span = date.fromisoformat(info["depart_date"]) - start
        arrive = time.fromisoformat(info["arrive_time"]) if info.get("arrive_time") else None
        depart = time.fromisoformat(info["depart_time"]) if info.get("depart_time") else None
    except (ValueError, KeyError, TypeError) as e:
        raise NotOurExport(DAMAGED) from e
    moved = arrive_date or start
    info |= {"arrive_date": moved.isoformat(), "depart_date": (moved + span).isoformat()}
    trip = Trip(arrive_date=moved, depart_date=moved + span, arrive_time=arrive, depart_time=depart)
    upload = Upload(trip=info)
    _rows(db, wb["Itinerary"], meta["rows"], trip, upload)
    return upload


def _rows(db: Session, ws, identities: dict, trip: Trip, upload: Upload) -> None:
    header = {str(c.value): c.column for c in ws[3] if c.value}
    col = {name: header.get(name) for name in ("Start", "End", "Block", "Details", ROW_KEY)}
    if not (col["Start"] and col["End"] and col["Block"]):
        raise NotOurExport("The Itinerary sheet's Start, End and Block columns are missing.")

    def cell(r: int, name: str):
        return ws.cell(row=r, column=col[name]).value if col[name] else None

    known = {i["place_id"] for i in identities.values() if i.get("place_id")}
    stored = (set(db.scalars(select(Place.place_id).where(Place.place_id.in_(known))))
              if known else set())
    seen_keys: set[str] = set()
    used_blocks: set[str] = set()
    limit = settings().max_stops_per_day
    day = None

    for r in range(4, ws.max_row + 1):
        band = BAND.match(str(ws.cell(row=r, column=1).value or ""))
        if band:
            day = int(band[1]) - 1
            continue
        title = str(cell(r, "Block") or "").strip()
        if day is None or not (title or cell(r, "Start") or cell(r, "End")):
            continue

        at, end = _clock(cell(r, "Start")), _end(cell(r, "End"), day)
        if not title:
            upload.skipped.append((r, "no title in Block"))
            continue
        if at is None or at >= DAY or end is None:
            upload.skipped.append((r, "Start or End is not a time like 14:30"))
            continue
        duration = (end[0] - day) * DAY + end[1] - at
        if duration <= 0:
            upload.skipped.append((r, "it ends before it starts"))
            continue

        key = str(cell(r, ROW_KEY) or "")
        identity = identities.get(key) if key not in seen_keys else None
        seen_keys.add(key)
        details = str(cell(r, "Details") or "").strip()[:DESCRIPTION_MAX] or None
        ref = identity.get("reference_url") if identity else None

        place_id = identity.get("place_id") if identity else None
        block_id = identity.get("block_id") if identity else None
        if not isinstance(block_id, str) or not 0 < len(block_id) <= 36 or block_id in used_blocks:
            block_id = str(uuid4())
        try:
            if place_id and title == identity.get("name") and place_id in stored:
                item = ItemIn(place_id=place_id, block_id=block_id, start_min=at,
                              duration_min=duration, description=details, reference_url=ref)
            else:
                item = ItemIn(kind=BlockKind.CUSTOM, block_id=block_id,
                              title=title[:TITLE_MAX], description=details, start_min=at,
                              duration_min=duration, reference_url=ref)
        except ValueError:
            upload.skipped.append((r, "its hidden Trip Planner data is damaged"))
            continue
        used_blocks.add(item.block_id)

        covered = pieces(day, at, duration)
        if day < 0 or covered[-1][0] >= day_count(trip):
            upload.skipped.append((r, f"Day {day + 1} is outside the trip"))
            continue
        if any(a < available_window(trip, d)[0] or b > available_window(trip, d)[1]
               for d, a, b in covered):
            upload.skipped.append((r, "it falls outside the flight times"))
            continue
        if len(upload.days.get(day, [])) >= limit:
            upload.skipped.append((r, f"Day {day + 1} already has {limit} blocks"))
            continue
        upload.days.setdefault(day, []).append(item)

