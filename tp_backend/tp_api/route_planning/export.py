"""The trip as one .xlsx: the ordered days, then the shortlist behind them.

Colours are tp_client/docs/design-system.md, and the warning English lives here because a
spreadsheet has no client to own it.
"""

import textwrap
from datetime import UTC, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from libs.db import Trip
from libs.routing import PlanWarning, Stop, hhmm, plan_day
from tp_api.deps import HoursLookup
from tp_api.route_planning import service
from tp_api.route_planning.utils import day_count, google_weekday

INK, PAPER, SAND = "252B20", "FDFBF3", "F3F0E6"
BRAND, BRAND_BG = "2C6B64", "E5EFEC"
OK, OK_BG = "4E7A45", "E9F0E0"
WARN, WARN_BG = "8A6524", "F6EBD4"
ALERT, ALERT_BG = "8B4526", "F8E7DC"
FAINT, HAIRLINE = "98A08D", "E3DDCB"

LINK_ICON = "🔗"

WARNING_TEXT = {
    "closed": "closed all day",
    "opens_later": "opens {opens}, you arrive {start}",
    "closes_before_done": "closes {closes} before you finish",
    "after_sunset": "dark by {start} — sunset {sunset}",
    "no_hours": "opening hours unknown",
}
BLOCKING = {"closed"}

ITINERARY_COLUMNS = [("Day", 6), ("Date", 10), ("#", 4), ("Start", 7), ("End", 7), ("Block", 32),
                     ("Category", 10), ("Warning", 36)]
DETAILS_COLUMN = ("Details", 44)
REFERENCE_COLUMN = ("Ref", 5)
SHORTLIST_COLUMNS = [("Place", 32), ("Category", 10), ("Mentions", 10), ("Used on", 10),
                     ("Why go", 60), ("Source", 7)]
CENTRED = {"#", "Start", "End", "Ref", "Source", "Mentions", "Used on"}
WRAPPED = {"Block", "Warning", "Details", "Why go"}

HAIRLINE_BOTTOM = Border(bottom=Side(style="thin", color=HAIRLINE))
LINE_HEIGHT = 15


def _header(ws: Worksheet, row: int, columns: list[tuple[str, int]]) -> None:
    for i, (label, width) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=i, value=label)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(bold=True, color=PAPER, size=11)
        cell.alignment = Alignment(horizontal="center" if label in CENTRED else "left")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _band(ws: Worksheet, row: int, columns: int, text: str) -> None:
    for i in range(1, columns + 1):
        ws.cell(row=row, column=i).fill = PatternFill("solid", fgColor=BRAND_BG)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True, color=BRAND)


def _lines(text, width: int) -> int:
    """Excel will not auto-fit a row that holds a merged cell, so the height is counted here."""
    return sum(len(textwrap.wrap(part, width)) or 1 for part in str(text).split("\n"))


def _body_row(ws: Worksheet, row: int, columns: list[tuple[str, int]], values: list,
              *, stripe: str, faint: set[int] = frozenset()) -> None:
    tallest = 1
    for i, (label, width) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=i, value=values[i - 1])
        cell.fill = PatternFill("solid", fgColor=stripe)
        cell.border = HAIRLINE_BOTTOM
        cell.font = Font(color=FAINT if i in faint else INK)
        wrap = label in WRAPPED
        cell.alignment = Alignment(horizontal="center" if label in CENTRED else "left",
                                   vertical="top", wrap_text=wrap)
        if wrap and values[i - 1]:
            tallest = max(tallest, _lines(values[i - 1], width))
    ws.row_dimensions[row].height = tallest * LINE_HEIGHT


def _merge_down(ws: Worksheet, top: int, bottom: int, columns: tuple[int, ...]) -> None:
    """Day and Date repeat unchanged down a band, so each becomes one tall cell."""
    if bottom <= top:
        return
    for column in columns:
        ws.merge_cells(start_row=top, start_column=column, end_row=bottom, end_column=column)
        ws.cell(row=top, column=column).alignment = Alignment(horizontal="center",
                                                             vertical="center")


def _link(ws: Worksheet, row: int, column: int, url: str | None, title: str | None) -> None:
    """One cell, one link icon — Excel and Sheets both open a hyperlink in the browser."""
    if not url:
        return
    cell = ws.cell(row=row, column=column, value=LINK_ICON)
    # Excel shows the hyperlink's own tooltip on hover, which is where the source title goes.
    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=url, tooltip=title)
    cell.font = Font(color=BRAND)
    cell.alignment = Alignment(horizontal="center")


def _tint(ws: Worksheet, row: int, column: int, fill: str, ink: str) -> None:
    cell = ws.cell(row=row, column=column)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(color=ink)


def _warning_cell(found: list[PlanWarning]) -> tuple[str, str, str]:
    text = "; ".join(WARNING_TEXT[w.code].format(**w.detail) for w in found) or "ok"
    if any(w.code in BLOCKING for w in found):
        return text, ALERT_BG, ALERT
    return (text, WARN_BG, WARN) if found else (text, OK_BG, OK)


def itinerary_sheet(ws: Worksheet, db: Session, trip: Trip, fetch: HoursLookup) -> None:
    city = trip.city
    rows = service.day_rows(db, trip.trip_id)
    details = any(r.ItineraryItem.description for r in rows)
    refs = any(r.ItineraryItem.reference_url for r in rows)
    columns = (ITINERARY_COLUMNS + ([DETAILS_COLUMN] if details else [])
               + ([REFERENCE_COLUMN] if refs else []))

    ws.cell(row=1, column=1, value=f"{trip.name or city.name} · "
            f"{trip.arrive_date:%d %b}–{trip.depart_date:%d %b %Y}").font = Font(
                bold=True, size=14, color=INK)
    ws.cell(row=2, column=1, value=f"exported {datetime.now(UTC):%d %b %Y} · opening hours were "
            "checked then, re-check nearer the date").font = Font(color=FAINT, size=10)
    _header(ws, 3, columns)

    place_ids = [r.Place.place_id for r in rows if r.Place]
    facts = service.mention_facts(db, place_ids)
    hours = service.load_hours(db, place_ids, fetch) if place_ids else {}

    by_day: dict[int, list] = {}
    for r in rows:
        by_day.setdefault(r.ItineraryItem.day_index, []).append(r)

    at = 4
    for day in range(day_count(trip)):
        day_date = trip.arrive_date + timedelta(days=day)
        items = by_day.get(day, [])
        venues = [r.Place for r in items if r.Place]

        sun = service.sun_by_place(db, venues, day_date, hours)
        sunrise, sunset = service.first_daylight(sun, [p.place_id for p in venues])
        daylight = (f" · daylight {hhmm(sunrise)}–{hhmm(sunset)}"
                    if sunrise is not None and sunset is not None else "")

        _band(ws, at, len(columns), f"Day {day + 1} · {day_date:%a %d %b} · "
              f"{len(items) or 'no'} block{'' if len(items) == 1 else 's'}{daylight}")
        at += 1

        plan = plan_day(
            [Stop(place_id=r.Place.place_id, name=r.Place.name,
                  category=service.category_of(r.Place, facts),
                  start_min=r.ItineraryItem.start_min, duration_min=r.ItineraryItem.duration_min,
                  periods=hours[r.Place.place_id].periods if r.Place.place_id in hours else None,
                  sunset_min=sun.get(r.Place.place_id, (None, None))[1])
             for r in items if r.Place],
            weekday=google_weekday(day_date))
        routed = {b.place_id: b for b in plan.blocks}

        found: dict[str, list[PlanWarning]] = {}
        for w in plan.warnings:
            found.setdefault(w.place_id or "", []).append(w)

        top = at
        for n, r in enumerate(items, start=1):
            item = r.ItineraryItem
            block = routed.get(item.place_id) if r.Place else None
            start, end = ((block.start_min, block.end_min) if block
                          else (item.start_min, item.start_min + item.duration_min))
            text, fill, ink = _warning_cell(found.get(item.place_id or "", []))
            _body_row(ws, at, columns,
                      [day + 1, f"{day_date:%d %b}", n, hhmm(start), hhmm(end),
                       r.Place.name if r.Place else item.title,
                       service.category_of(r.Place, facts) or "" if r.Place else "",
                       text if block else None]
                      + ([item.description] if details else [])
                      + ([None] if refs else []),
                      stripe=SAND if n % 2 else PAPER, faint={4, 5})
            if block:
                _tint(ws, at, 8, fill, ink)
            if refs:
                _link(ws, at, len(columns), item.reference_url, "Your link for this block")
            at += 1
        _merge_down(ws, top, at - 1, (1, 2))


def shortlist_sheet(ws: Worksheet, db: Session, trip: Trip) -> None:
    """What the videos and notes offered, so the places left out are still in the file."""
    columns = SHORTLIST_COLUMNS
    ws.cell(row=1, column=1, value="Shortlist · ranked by how many sources named it").font = Font(
        bold=True, size=14, color=INK)
    _header(ws, 3, columns)

    listing = service.shortlist(db, trip.trip_id, limit=200, offset=0, category=None)
    for n, place in enumerate(listing.places, start=1):
        at = 3 + n
        used = place.day_index is not None
        src = next(iter(place.sources), None)
        _body_row(ws, at, columns,
                  [place.name, place.category or "", place.mention_count,
                   f"Day {place.day_index + 1}" if used else "not used", place.why_go or "", None],
                  stripe=SAND if n % 2 else PAPER, faint=set() if used else {4})
        if used:
            _tint(ws, at, 4, BRAND_BG, BRAND)
        _link(ws, at, 6, src.url if src else None, src.title if src else None)


def workbook_bytes(db: Session, trip: Trip, fetch: HoursLookup) -> bytes:
    wb = Workbook()
    itinerary_sheet(wb.active, db, trip, fetch)
    wb.active.title = "Itinerary"
    shortlist_sheet(wb.create_sheet("Shortlist"), db, trip)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(trip: Trip) -> str:
    stem = (trip.name or trip.city.name).replace(" ", "-")
    return f"{stem}-{trip.arrive_date:%Y-%m-%d}.xlsx"
